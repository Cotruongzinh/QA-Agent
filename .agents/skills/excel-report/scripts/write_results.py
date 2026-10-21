"""Validate qa-result.json and upsert a copied QA workbook report.

The JSON document is validated in full before any output directory, workbook
copy, or workbook write is attempted.  The source template is never opened as
an editable workbook: it is copied with shutil.copy2 and only the copy is
loaded, changed, saved, and reopened for verification.
"""

from __future__ import annotations

import argparse
import copy
import hashlib
import json
import math
import shutil
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Mapping, Sequence

from openpyxl import load_workbook
from openpyxl.drawing.image import Image as OpenpyxlImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.utils.units import pixels_to_EMU
from openpyxl.worksheet.worksheet import Worksheet
from PIL import Image as PillowImage


DEFAULT_INPUT = Path(r"C:\QA-Agent\runs\fail-smoke\qa-result.json")
DEFAULT_TEMPLATE = Path(r"C:\QA-Agent\templates\QA_Report_Template.xlsx")
DEFAULT_OUTPUT = Path(r"C:\QA-Agent\output\QA_Report_FAIL_SMOKE.xlsx")
DEFAULT_TEMPLATES_DIR = Path(r"C:\QA-Agent\templates")

TESTCASE_SHEET = "Testcase"
BUG_SHEET = "NoteBug"
SUMMARY_SHEET = "Summary"
ALLOWED_MODES = {"EXECUTE", "DISCOVER"}
ALLOWED_STATUSES = {"PASS", "FAIL", "UNVERIFIED"}
ALLOWED_EVIDENCE_TYPES = {"screenshot", "network", "console", "url", "observation"}

FIELD_TO_HEADER = {
    "tc_id": "TC_ID",
    "module": "Module",
    "test_title": "Test Title",
    "description": "Description/Test Summary",
    "test_steps": "Test Steps",
    "test_data": "Test Data",
    "expected_result": "Expected Result",
    "actual_result": "Actual Result",
    "evidence": "Log/Screenshot/Video",
    "status": "Result",
}

BUG_FIELD_TO_HEADER = {
    "bug_id": "Bug ID",
    "reporter": "Người report",
    "device": "Thiết bị",
    "os_app_version": "OS/App version",
    "module": "Module",
    "linked_tc_id": "Testcase liên quan",
    "title": "Tiêu đề bug",
    "reproduction_steps": "Bước tái hiện",
    "actual_result": "Kết quả thực tế",
    "expected_result": "Kết quả mong muốn",
    "evidence": "Log/Screenshot/Video",
}

RUN_FIELDS = ("run_id", "target", "mode", "browser", "started_at", "finished_at")
TESTCASE_FIELDS = (
    "tc_id", "module", "test_title", "description", "test_steps", "test_data",
    "expected_result", "actual_result", "severity", "status", "evidence", "notes",
)
BUG_FIELDS = (
    "bug_id", "linked_tc_id", "module", "title", "severity", "reproduction_steps",
    "actual_result", "expected_result", "evidence",
)
SUMMARY_FIELDS = ("total", "pass", "fail", "unverified", "bugs")
STYLE_ATTRIBUTES = ("font", "fill", "border", "alignment", "number_format", "protection")

SUMMARY_LABELS = {
    "total": {
        "tổng testcase", "tổng test case", "tổng số testcase", "tổng số test case",
        "total", "total testcase", "total testcases",
    },
    "pass": {"pass", "passed"},
    "fail": {"fail", "failed"},
    "unverified": {
        "unverified", "chưa kiểm thử", "chưa verify", "chưa xác minh", "not run",
    },
    "bugs": {
        "bug riêng ghi nhận", "bug ghi nhận", "tổng bug", "số bug", "bug count", "bugs",
    },
    "pass_rate": {"tỷ lệ pass", "tỉ lệ pass", "pass rate"},
}
REQUIRED_SUMMARY_KPIS = ("total", "pass", "fail", "bugs", "pass_rate")

MAX_IMAGE_WIDTH_PX = 320
MAX_IMAGE_HEIGHT_PX = 220
IMAGE_GAP_PX = 8
EVIDENCE_TEXT_LINE_HEIGHT_PX = 18


class ReportError(RuntimeError):
    """Raised for validation, safety, write, or verification failures."""


@dataclass(frozen=True)
class ValidatedInput:
    path: Path
    document: dict[str, Any]


@dataclass(frozen=True)
class WorkbookSnapshot:
    sheet_names: tuple[str, ...]
    image_counts: dict[str, int]
    image_signatures: dict[str, tuple[tuple[Any, ...], ...]]
    merged_counts: dict[str, int]
    merged_ranges: dict[str, tuple[str, ...]]
    formula_counts: dict[str, int]
    formula_states: dict[str, tuple[tuple[str, str], ...]]
    summary_state: dict[str, Any] | None


@dataclass(frozen=True)
class WriteRecord:
    row: int
    tc_id: str
    status: str
    expected_result: str
    actual_result: str
    evidence_text: str
    evidence_handling: str
    screenshot_count: int
    prior_screenshot_count: int
    operation: str
    definition_values: tuple[tuple[str, Any], ...]
    feedback_values: tuple[tuple[str, Any], ...]


@dataclass(frozen=True)
class BugWriteRecord:
    row: int
    bug_id: str
    linked_tc_id: str
    actual_result: str
    expected_result: str
    reporter: str
    device: str
    os_app_version: str
    module: str
    title: str
    reproduction_steps: str
    evidence_text: str
    evidence_handling: str
    screenshot_count: int
    prior_screenshot_count: int
    operation: str


@dataclass(frozen=True)
class SummaryMetrics:
    total: int
    passed: int
    failed: int
    unverified: int
    bugs: int
    pass_rate: float


@dataclass(frozen=True)
class SummaryWriteRecord:
    metrics: SummaryMetrics
    cells: dict[str, str]
    formula_replacements: frozenset[str]


def normalized(value: object) -> str:
    if value is None:
        return ""
    return " ".join(str(value).strip().casefold().split())


def verification_values_equal(actual: object, expected: object) -> bool:
    """Treat only None and the empty string as the same verification value."""
    if actual in (None, "") and expected in (None, ""):
        return True
    return actual == expected


def resolved(path: Path) -> Path:
    return path.expanduser().resolve(strict=False)


def is_within(path: Path, directory: Path) -> bool:
    try:
        path.relative_to(directory)
        return True
    except ValueError:
        return False


def is_in_templates_directory(path: Path) -> bool:
    if is_within(path, resolved(DEFAULT_TEMPLATES_DIR)):
        return True
    return any(parent.name.casefold() == "templates" for parent in path.parents)


def require_object(value: Any, location: str, errors: list[str]) -> dict[str, Any] | None:
    if not isinstance(value, dict):
        errors.append(f"{location} must be an object")
        return None
    return value


def require_fields(item: Mapping[str, Any], fields: Sequence[str], location: str, errors: list[str]) -> None:
    for field in fields:
        if field not in item:
            errors.append(f"{location}.{field} is required")


def validate_string_fields(
    item: Mapping[str, Any], fields: Sequence[str], location: str, errors: list[str]
) -> None:
    for field in fields:
        if field in item and not isinstance(item[field], str):
            errors.append(f"{location}.{field} must be a string")


def validate_string_list(value: Any, location: str, errors: list[str]) -> None:
    if not isinstance(value, list):
        errors.append(f"{location} must be a list")
        return
    for index, entry in enumerate(value):
        if not isinstance(entry, str):
            errors.append(f"{location}[{index}] must be a string")


def validate_local_evidence_path(path_text: Any, location: str, errors: list[str]) -> Path | None:
    if not isinstance(path_text, str) or not path_text.strip():
        errors.append(f"{location} must be a non-empty local path string")
        return None
    if path_text.strip().casefold().startswith(("http://", "https://", "data:")):
        errors.append(f"{location} must be a local filesystem path, not a URL/base64 value")
        return None
    evidence_path = Path(path_text)
    if not evidence_path.is_absolute():
        errors.append(f"{location} must be an absolute local filesystem path")
        return None
    evidence_path = resolved(evidence_path)
    if not evidence_path.is_file():
        errors.append(f"{location} does not exist or is not a file: {evidence_path}")
        return None
    return evidence_path


def validate_screenshot(path: Path, location: str, errors: list[str]) -> None:
    try:
        with PillowImage.open(path) as image:
            image.verify()
        # openpyxl performs its own Pillow-backed format check here.
        OpenpyxlImage(str(path))
    except Exception as exc:
        errors.append(f"{location} is not a Pillow/openpyxl-supported image: {exc}")


def validate_evidence(value: Any, location: str, errors: list[str]) -> None:
    if not isinstance(value, list):
        errors.append(f"{location} must be a list")
        return
    for index, raw_entry in enumerate(value):
        entry_location = f"{location}[{index}]"
        entry = require_object(raw_entry, entry_location, errors)
        if entry is None:
            continue
        require_fields(entry, ("type", "path", "description"), entry_location, errors)
        validate_string_fields(entry, ("type", "path", "description"), entry_location, errors)
        evidence_type = entry.get("type")
        if isinstance(evidence_type, str) and evidence_type not in ALLOWED_EVIDENCE_TYPES:
            errors.append(
                f"{entry_location}.type must be one of {sorted(ALLOWED_EVIDENCE_TYPES)}"
            )
        evidence_path = validate_local_evidence_path(entry.get("path"), f"{entry_location}.path", errors)
        if evidence_type == "screenshot" and evidence_path is not None:
            validate_screenshot(evidence_path, f"{entry_location}.path", errors)


def load_and_validate_json(input_path: Path) -> ValidatedInput:
    path = resolved(input_path)
    errors: list[str] = []
    if not path.is_file():
        raise ReportError(f"Input JSON not found: {path}")
    try:
        with path.open("r", encoding="utf-8-sig") as stream:
            raw_document = json.load(stream)
    except (OSError, json.JSONDecodeError) as exc:
        raise ReportError(f"Cannot read valid JSON from {path}: {exc}") from exc

    document = require_object(raw_document, "$", errors)
    if document is None:
        raise ReportError("qa-result.json validation failed:\n- " + "\n- ".join(errors))
    require_fields(document, ("run", "testcases", "bugs", "summary"), "$", errors)

    run = require_object(document.get("run"), "$.run", errors)
    if run is not None:
        require_fields(run, RUN_FIELDS, "$.run", errors)
        validate_string_fields(run, RUN_FIELDS, "$.run", errors)
        if run.get("mode") not in ALLOWED_MODES:
            errors.append(f"$.run.mode must be one of {sorted(ALLOWED_MODES)}")

    testcases = document.get("testcases")
    testcase_ids: list[str] = []
    if not isinstance(testcases, list):
        errors.append("$.testcases must be a list")
        testcases = []
    else:
        for index, raw_testcase in enumerate(testcases):
            location = f"$.testcases[{index}]"
            testcase = require_object(raw_testcase, location, errors)
            if testcase is None:
                continue
            require_fields(testcase, TESTCASE_FIELDS, location, errors)
            validate_string_fields(
                testcase,
                ("tc_id", "module", "test_title", "description", "test_data",
                 "expected_result", "actual_result", "severity", "status", "notes"),
                location,
                errors,
            )
            tc_id = testcase.get("tc_id")
            if isinstance(tc_id, str):
                if not tc_id.strip():
                    errors.append(f"{location}.tc_id must not be empty")
                else:
                    testcase_ids.append(tc_id)
            status = testcase.get("status")
            if isinstance(status, str) and status not in ALLOWED_STATUSES:
                errors.append(f"{location}.status must be one of {sorted(ALLOWED_STATUSES)}")
            validate_string_list(testcase.get("test_steps"), f"{location}.test_steps", errors)
            validate_evidence(testcase.get("evidence"), f"{location}.evidence", errors)

    duplicates = sorted({tc_id for tc_id in testcase_ids if testcase_ids.count(tc_id) > 1})
    if duplicates:
        errors.append("duplicate testcase tc_id value(s): " + ", ".join(duplicates))

    bugs = document.get("bugs")
    bug_ids: list[str] = []
    if not isinstance(bugs, list):
        errors.append("$.bugs must be a list")
        bugs = []
    else:
        known_ids = set(testcase_ids)
        for index, raw_bug in enumerate(bugs):
            location = f"$.bugs[{index}]"
            bug = require_object(raw_bug, location, errors)
            if bug is None:
                continue
            require_fields(bug, BUG_FIELDS, location, errors)
            validate_string_fields(
                bug,
                ("bug_id", "linked_tc_id", "module", "title", "severity",
                 "actual_result", "expected_result"),
                location,
                errors,
            )
            validate_string_list(bug.get("reproduction_steps"), f"{location}.reproduction_steps", errors)
            validate_evidence(bug.get("evidence"), f"{location}.evidence", errors)
            bug_id = bug.get("bug_id")
            if isinstance(bug_id, str):
                if not bug_id.strip():
                    errors.append(f"{location}.bug_id must not be empty")
                else:
                    bug_ids.append(bug_id)
            linked_tc_id = bug.get("linked_tc_id")
            if isinstance(linked_tc_id, str) and linked_tc_id not in known_ids:
                errors.append(f"{location}.linked_tc_id does not reference an existing testcase: {linked_tc_id!r}")

    duplicate_bug_ids = sorted({bug_id for bug_id in bug_ids if bug_ids.count(bug_id) > 1})
    if duplicate_bug_ids:
        errors.append("duplicate bug_id value(s): " + ", ".join(duplicate_bug_ids))

    if run is not None and (bugs or "environment" in run):
        environment = require_object(run.get("environment"), "$.run.environment", errors)
        if environment is not None:
            environment_fields = ("reporter", "device", "os_app_version")
            require_fields(environment, environment_fields, "$.run.environment", errors)
            validate_string_fields(environment, environment_fields, "$.run.environment", errors)
            for field in environment_fields:
                value = environment.get(field)
                if isinstance(value, str) and not value.strip():
                    errors.append(
                        f"$.run.environment.{field} must not be empty; use NOT_RECORDED when unknown"
                    )

    summary = require_object(document.get("summary"), "$.summary", errors)
    if summary is not None:
        require_fields(summary, SUMMARY_FIELDS, "$.summary", errors)
        for field in SUMMARY_FIELDS:
            value = summary.get(field)
            if not isinstance(value, int) or isinstance(value, bool) or value < 0:
                errors.append(f"$.summary.{field} must be a non-negative integer")
        if all(isinstance(summary.get(field), int) and not isinstance(summary.get(field), bool) for field in SUMMARY_FIELDS):
            status_counts = {
                status: sum(1 for testcase in testcases if isinstance(testcase, dict) and testcase.get("status") == status)
                for status in ALLOWED_STATUSES
            }
            expected_counts = {
                "total": len(testcases),
                "pass": status_counts["PASS"],
                "fail": status_counts["FAIL"],
                "unverified": status_counts["UNVERIFIED"],
                "bugs": len(bugs),
            }
            for field, expected in expected_counts.items():
                if summary.get(field) != expected:
                    errors.append(
                        f"$.summary.{field}={summary.get(field)!r} does not match expected {expected}"
                    )
            if summary.get("pass", 0) + summary.get("fail", 0) + summary.get("unverified", 0) != summary.get("total"):
                errors.append("$.summary pass + fail + unverified must equal total")

    if errors:
        raise ReportError("qa-result.json validation failed:\n- " + "\n- ".join(errors))
    return ValidatedInput(path=path, document=document)


def validate_paths(template_path: Path, output_path: Path) -> tuple[Path, Path]:
    template = resolved(template_path)
    output = resolved(output_path)
    if template == output:
        raise ReportError("Refusing to write: template and output paths are identical")
    if is_in_templates_directory(output):
        raise ReportError("Refusing to write: output is inside a templates directory")
    if not template.is_file():
        raise ReportError(f"Template not found: {template}")
    if template.suffix.casefold() != ".xlsx" or output.suffix.casefold() != ".xlsx":
        raise ReportError("Template and output must both be .xlsx files")
    return template, output


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def is_formula(value: object) -> bool:
    return isinstance(value, str) and value.startswith("=")


def formula_count(ws: Worksheet) -> int:
    return sum(
        is_formula(cell.value)
        for row in ws.iter_rows()
        for cell in row
        if cell.value is not None
    )


def formula_state(ws: Worksheet) -> tuple[tuple[str, str], ...]:
    return tuple(
        (cell.coordinate, cell.value)
        for row in ws.iter_rows()
        for cell in row
        if is_formula(cell.value)
    )


def image_anchor_signature(image: object) -> tuple[Any, ...]:
    anchor = getattr(image, "anchor", None)
    marker = getattr(anchor, "_from", None)
    return (
        getattr(marker, "row", None), getattr(marker, "col", None),
        getattr(marker, "rowOff", None), getattr(marker, "colOff", None),
        getattr(image, "width", None), getattr(image, "height", None),
    )


def color_semantics(color: object) -> tuple[Any, ...] | None:
    if color is None:
        return None
    color_type = getattr(color, "type", None)
    tint = getattr(color, "tint", 0.0) or 0.0
    if color_type == "rgb":
        return ("rgb", getattr(color, "rgb", None), tint)
    if color_type == "indexed":
        return ("indexed", getattr(color, "indexed", None), tint)
    if color_type == "theme":
        return ("theme", getattr(color, "theme", None), tint)
    if color_type == "auto":
        return ("auto", bool(getattr(color, "auto", False)), tint)
    return (color_type, tint)


def side_semantics(side: object) -> tuple[Any, ...] | None:
    if side is None:
        return None
    return (getattr(side, "style", None), color_semantics(getattr(side, "color", None)))


def font_semantics(font: object) -> tuple[Any, ...]:
    underline = getattr(font, "u", None)
    if underline == "none":
        underline = None
    return (
        getattr(font, "name", None), getattr(font, "sz", None),
        bool(getattr(font, "b", False)), bool(getattr(font, "i", False)),
        getattr(font, "vertAlign", None), underline,
        bool(getattr(font, "strike", False)), color_semantics(getattr(font, "color", None)),
        bool(getattr(font, "outline", False)), bool(getattr(font, "shadow", False)),
        bool(getattr(font, "condense", False)), bool(getattr(font, "extend", False)),
        getattr(font, "scheme", None), getattr(font, "family", None),
        getattr(font, "charset", None),
    )


def fill_semantics(fill: object) -> tuple[Any, ...]:
    stops = tuple(
        (getattr(stop, "position", None), color_semantics(getattr(stop, "color", None)))
        for stop in (getattr(fill, "stop", None) or ())
    )
    fill_type = getattr(fill, "fill_type", None)
    if fill_type in ("linear", "path"):
        return (
            fill_type, getattr(fill, "degree", None), getattr(fill, "left", None),
            getattr(fill, "right", None), getattr(fill, "top", None),
            getattr(fill, "bottom", None), stops,
        )
    return (
        fill_type,
        color_semantics(getattr(fill, "fgColor", None)),
        color_semantics(getattr(fill, "bgColor", None)),
    )


def border_semantics(border: object) -> tuple[Any, ...]:
    return (
        side_semantics(getattr(border, "left", None)),
        side_semantics(getattr(border, "right", None)),
        side_semantics(getattr(border, "top", None)),
        side_semantics(getattr(border, "bottom", None)),
        side_semantics(getattr(border, "diagonal", None)),
        side_semantics(getattr(border, "vertical", None)),
        side_semantics(getattr(border, "horizontal", None)),
        bool(getattr(border, "diagonalUp", False)),
        bool(getattr(border, "diagonalDown", False)),
        True if getattr(border, "outline", None) is None else bool(border.outline),
    )


def alignment_semantics(alignment: object) -> tuple[Any, ...]:
    return (
        getattr(alignment, "horizontal", None), getattr(alignment, "vertical", None),
        getattr(alignment, "text_rotation", 0) or 0,
        bool(getattr(alignment, "wrap_text", False)),
        bool(getattr(alignment, "shrink_to_fit", False)),
        getattr(alignment, "indent", 0) or 0,
        getattr(alignment, "relativeIndent", 0) or 0,
        bool(getattr(alignment, "justifyLastLine", False)),
        getattr(alignment, "readingOrder", 0) or 0,
    )


def protection_semantics(protection: object) -> tuple[Any, ...]:
    locked = getattr(protection, "locked", None)
    hidden = getattr(protection, "hidden", None)
    return (True if locked is None else bool(locked), bool(hidden))


def cell_semantics(cell: object) -> dict[str, Any]:
    return {
        "value/formula": cell.value,
        "font": font_semantics(cell.font),
        "fill": fill_semantics(cell.fill),
        "border": border_semantics(cell.border),
        "alignment": alignment_semantics(cell.alignment),
        "number_format": cell.number_format,
        "protection": protection_semantics(cell.protection),
    }


def summary_semantic_state(ws: Worksheet) -> dict[str, Any]:
    return {
        "cells": {
            cell.coordinate: cell_semantics(cell)
            for row in ws.iter_rows()
            for cell in row
            if cell.value is not None or cell.has_style
        },
        "merged_ranges": tuple(sorted(str(cell_range) for cell_range in ws.merged_cells.ranges)),
        "row_heights": {
            index: dimension.height for index, dimension in sorted(ws.row_dimensions.items())
        },
        "column_widths": {
            key: dimension.width for key, dimension in sorted(ws.column_dimensions.items())
        },
        "hidden_rows": {
            index: bool(dimension.hidden) for index, dimension in sorted(ws.row_dimensions.items())
        },
        "hidden_columns": {
            key: bool(dimension.hidden) for key, dimension in sorted(ws.column_dimensions.items())
        },
        "image_count": len(getattr(ws, "_images", ())),
        "freeze_panes": str(ws.freeze_panes or ""),
    }


def summary_semantic_differences(
    before: dict[str, Any] | None,
    after: dict[str, Any] | None,
    allowed_value_cells: set[str],
) -> list[str]:
    if before is None or after is None:
        return [] if before is after else ["Summary sheet presence changed"]
    differences: list[str] = []
    before_cells = before["cells"]
    after_cells = after["cells"]
    for coordinate in sorted(set(before_cells) | set(after_cells)):
        old = before_cells.get(coordinate)
        new = after_cells.get(coordinate)
        if old is None or new is None:
            if coordinate in allowed_value_cells:
                continue
            differences.append(f"Summary!{coordinate} cell presence changed")
            continue
        for property_name in old:
            if old[property_name] == new[property_name]:
                continue
            if property_name == "value/formula" and coordinate in allowed_value_cells:
                continue
            differences.append(
                f"Summary!{coordinate} {property_name} changed: "
                f"before={old[property_name]!r}, after={new[property_name]!r}"
            )
    old_merged = set(before["merged_ranges"])
    new_merged = set(after["merged_ranges"])
    for cell_range in sorted(old_merged - new_merged):
        differences.append(f"Summary merged range removed: {cell_range}")
    for cell_range in sorted(new_merged - old_merged):
        differences.append(f"Summary merged range added: {cell_range}")
    mapping_properties = (
        ("row_heights", "row", "height"),
        ("column_widths", "column", "width"),
        ("hidden_rows", "row", "hidden"),
        ("hidden_columns", "column", "hidden"),
    )
    for property_name, item_kind, display_property in mapping_properties:
        old_mapping = before[property_name]
        new_mapping = after[property_name]
        for key in sorted(set(old_mapping) | set(new_mapping)):
            if old_mapping.get(key) != new_mapping.get(key):
                differences.append(
                    f"Summary {item_kind} {key} {display_property} changed: "
                    f"before={old_mapping.get(key)!r}, after={new_mapping.get(key)!r}"
                )
    for property_name in ("image_count", "freeze_panes"):
        if before[property_name] != after[property_name]:
            differences.append(
                f"Summary {property_name} changed: "
                f"before={before[property_name]!r}, after={after[property_name]!r}"
            )
    return differences


def snapshot(workbook: object) -> WorkbookSnapshot:
    return WorkbookSnapshot(
        sheet_names=tuple(workbook.sheetnames),
        image_counts={ws.title: len(getattr(ws, "_images", ())) for ws in workbook.worksheets},
        image_signatures={
            ws.title: tuple(image_anchor_signature(image) for image in getattr(ws, "_images", ()))
            for ws in workbook.worksheets
        },
        merged_counts={ws.title: len(ws.merged_cells.ranges) for ws in workbook.worksheets},
        merged_ranges={
            ws.title: tuple(sorted(str(cell_range) for cell_range in ws.merged_cells.ranges))
            for ws in workbook.worksheets
        },
        formula_counts={ws.title: formula_count(ws) for ws in workbook.worksheets},
        formula_states={ws.title: formula_state(ws) for ws in workbook.worksheets},
        summary_state=(
            summary_semantic_state(workbook[SUMMARY_SHEET])
            if SUMMARY_SHEET in workbook.sheetnames else None
        ),
    )


def find_last_header(ws: Worksheet, header: str) -> tuple[int, int]:
    occurrences = [
        (cell.row, cell.column)
        for row in ws.iter_rows()
        for cell in row
        if normalized(cell.value) == normalized(header)
    ]
    if not occurrences:
        raise ReportError(f"{header} header not found on sheet {ws.title}")
    return occurrences[-1]


def header_columns(
    ws: Worksheet, header_row: int, field_to_header: Mapping[str, str]
) -> dict[str, int]:
    columns: dict[str, int] = {}
    duplicates: set[str] = set()
    for column in range(1, ws.max_column + 1):
        label = normalized(ws.cell(header_row, column).value)
        if not label:
            continue
        if label in columns:
            duplicates.add(label)
        else:
            columns[label] = column
    required = {normalized(header): header for header in field_to_header.values()}
    missing = [display for key, display in required.items() if key not in columns]
    ambiguous = [required[key] for key in duplicates if key in required]
    if missing:
        raise ReportError(f"Missing {ws.title} header(s): " + ", ".join(missing))
    if ambiguous:
        raise ReportError(f"Duplicate {ws.title} header(s): " + ", ".join(ambiguous))
    return columns


def current_identifier_rows(
    ws: Worksheet, header_row: int, identifier_column: int, identifier_header: str
) -> list[int]:
    return [
        row
        for row in range(header_row + 1, ws.max_row + 1)
        if normalized(ws.cell(row, identifier_column).value)
        and normalized(ws.cell(row, identifier_column).value) != normalized(identifier_header)
    ]


def unique_identifier_rows(
    ws: Worksheet, header_row: int, identifier_column: int, identifier_header: str
) -> dict[str, int]:
    rows_by_identifier: dict[str, list[int]] = {}
    for row in current_identifier_rows(ws, header_row, identifier_column, identifier_header):
        identifier = str(ws.cell(row, identifier_column).value)
        rows_by_identifier.setdefault(identifier, []).append(row)
    duplicates = {
        identifier: rows for identifier, rows in rows_by_identifier.items() if len(rows) > 1
    }
    if duplicates:
        details = ", ".join(
            f"{identifier!r} at rows {rows}" for identifier, rows in sorted(duplicates.items())
        )
        raise ReportError(f"Duplicate {identifier_header} value(s) in current {ws.title} block: {details}")
    return {identifier: rows[0] for identifier, rows in rows_by_identifier.items()}


def row_is_blank(ws: Worksheet, row: int, mapped_columns: Sequence[int]) -> bool:
    return all(ws.cell(row, column).value in (None, "") for column in mapped_columns)


def row_has_anchored_image(ws: Worksheet, row: int) -> bool:
    return any(
        getattr(getattr(image, "anchor", None), "_from", None) is not None
        and image.anchor._from.row == row - 1
        for image in getattr(ws, "_images", ())
    )


def target_rows(
    ws: Worksheet,
    header_row: int,
    identifier_column: int,
    identifier_header: str,
    mapped_columns: Sequence[int],
    count: int,
) -> tuple[list[int], int | None]:
    existing = current_identifier_rows(
        ws, header_row, identifier_column, identifier_header
    )
    last_data_row = max(existing) if existing else None
    candidate = (last_data_row + 1) if last_data_row is not None else (header_row + 1)
    targets: list[int] = []
    while len(targets) < count:
        if row_is_blank(ws, candidate, mapped_columns) and not row_has_anchored_image(ws, candidate):
            targets.append(candidate)
        candidate += 1
    return targets, last_data_row


def next_append_row(
    ws: Worksheet,
    header_row: int,
    identifier_column: int,
    identifier_header: str,
    mapped_columns: Sequence[int],
) -> tuple[int, int | None]:
    rows, last_data_row = target_rows(
        ws, header_row, identifier_column, identifier_header, mapped_columns, 1
    )
    return rows[0], last_data_row


def anchored_image_count(ws: Worksheet, row: int, column: int) -> int:
    return sum(
        image_anchor_signature(image)[:2] == (row - 1, column - 1)
        for image in getattr(ws, "_images", ())
    )


def copy_cell_style(source: object, target: object) -> None:
    for attribute in STYLE_ATTRIBUTES:
        setattr(target, attribute, copy.copy(getattr(source, attribute)))


def nearest_style_row(
    ws: Worksheet, header_row: int, target_row: int, last_testcase_row: int | None,
    mapped_columns: Sequence[int],
) -> int:
    start_row = min(last_testcase_row or (target_row - 1), target_row - 1)
    for row in range(start_row, header_row, -1):
        if any(ws.cell(row, column).has_style for column in mapped_columns):
            return row
    for row in range(header_row + 1, max(ws.max_row, target_row) + 1):
        if any(ws.cell(row, column).has_style for column in mapped_columns):
            return row
    raise ReportError(f"No data-row style is available to copy on sheet {ws.title}")


def prepare_row_style(
    ws: Worksheet, source_row: int, target_row: int, mapped_columns: Sequence[int]
) -> None:
    for column in mapped_columns:
        copy_cell_style(ws.cell(source_row, column), ws.cell(target_row, column))
    source_height = ws.row_dimensions[source_row].height
    if source_height is not None:
        ws.row_dimensions[target_row].height = source_height


def prepare_row_style_if_unstyled(
    ws: Worksheet, source_row: int, target_row: int, mapped_columns: Sequence[int]
) -> None:
    if not any(ws.cell(target_row, column).has_style for column in mapped_columns):
        prepare_row_style(ws, source_row, target_row, mapped_columns)


def find_status_style_row(
    ws: Worksheet, header_row: int, target_row: int, tc_id_column: int,
    result_column: int, status: str,
) -> int | None:
    for row in range(target_row - 1, header_row, -1):
        if normalized(ws.cell(row, tc_id_column).value) and normalized(ws.cell(row, result_column).value) == normalized(status):
            return row
    for row in range(target_row + 1, ws.max_row + 1):
        if normalized(ws.cell(row, tc_id_column).value) and normalized(ws.cell(row, result_column).value) == normalized(status):
            return row
    return None


def format_steps(steps: Sequence[str]) -> str:
    return "\n".join(f"{index}. {step}" for index, step in enumerate(steps, start=1))


def evidence_text(evidence: Sequence[Mapping[str, str]]) -> str:
    lines: list[str] = []
    for item in evidence:
        if item["type"] == "screenshot":
            continue
        description = item["description"]
        path = resolved(Path(item["path"]))
        lines.append(f"{description}\n{path}" if description else str(path))
    return "\n".join(line for line in lines if line)


def add_screenshots(
    ws: Worksheet, row: int, column: int, evidence: Sequence[Mapping[str, str]], text_lines: int
) -> int:
    screenshots = [item for item in evidence if item["type"] == "screenshot"]
    if not screenshots:
        return 0
    y_offset = max(text_lines, 1) * EVIDENCE_TEXT_LINE_HEIGHT_PX
    for item in screenshots:
        image = OpenpyxlImage(str(resolved(Path(item["path"]))))
        scale = min(1.0, MAX_IMAGE_WIDTH_PX / image.width, MAX_IMAGE_HEIGHT_PX / image.height)
        image.width = int(image.width * scale)
        image.height = int(image.height * scale)
        marker = AnchorMarker(
            col=column - 1,
            row=row - 1,
            colOff=0,
            rowOff=pixels_to_EMU(y_offset),
        )
        image.anchor = OneCellAnchor(
            _from=marker,
            ext=XDRPositiveSize2D(pixels_to_EMU(image.width), pixels_to_EMU(image.height)),
        )
        ws.add_image(image)
        y_offset += image.height + IMAGE_GAP_PX
    needed_height = y_offset * 0.75
    current_height = ws.row_dimensions[row].height or 15.0
    ws.row_dimensions[row].height = max(current_height, needed_height)
    return len(screenshots)


def write_testcases(
    workbook: object, validated: ValidatedInput
) -> tuple[list[WriteRecord], int, dict[str, int]]:
    if TESTCASE_SHEET not in workbook.sheetnames:
        raise ReportError("Required sheet not found: Testcase")
    ws = workbook[TESTCASE_SHEET]
    header_row, tc_id_column = find_last_header(ws, "TC_ID")
    headers = header_columns(ws, header_row, FIELD_TO_HEADER)
    mapped_columns = sorted({headers[normalized(header)] for header in FIELD_TO_HEADER.values()})
    occupancy_columns = list(mapped_columns)
    for optional_header in ("Phản hồi", "FEEDBACK"):
        optional_column = headers.get(normalized(optional_header))
        if optional_column is not None:
            occupancy_columns.append(optional_column)
    testcases: list[dict[str, Any]] = validated.document["testcases"]
    existing_ids = unique_identifier_rows(ws, header_row, tc_id_column, "TC_ID")
    records: list[WriteRecord] = []
    result_column = headers[normalized("Result")]
    evidence_column = headers[normalized("Log/Screenshot/Video")]
    definition_fields = (
        "tc_id", "module", "test_title", "description", "test_steps", "test_data",
        "expected_result",
    )
    for testcase in testcases:
        existing_row = existing_ids.get(testcase["tc_id"])
        operation = "updated existing" if existing_row is not None else "appended new"
        if existing_row is not None:
            row = existing_row
            workbook_expected = ws.cell(
                row, headers[normalized("Expected Result")]
            ).value
            if not verification_values_equal(workbook_expected, testcase["expected_result"]):
                raise ReportError(
                    f"Definition mismatch for {testcase['tc_id']} Expected Result: "
                    f"workbook={workbook_expected!r}, JSON={testcase['expected_result']!r}"
                )
        else:
            row, last_testcase_row = next_append_row(
                ws, header_row, tc_id_column, "TC_ID", occupancy_columns
            )
            if not row_is_blank(ws, row, occupancy_columns):
                raise ReportError(f"Target row is not blank in mapped columns: Testcase!{row}")
            source_row = nearest_style_row(
                ws, header_row, row, last_testcase_row, mapped_columns
            )
            prepare_row_style_if_unstyled(ws, source_row, row, mapped_columns)

        current_status_matches = (
            existing_row is not None
            and normalized(ws.cell(row, result_column).value) == normalized(testcase["status"])
        )
        status_style_row = row if current_status_matches else find_status_style_row(
            ws, header_row, row, tc_id_column, result_column, testcase["status"]
        )
        if testcase["status"] == "PASS" and status_style_row is None:
            raise ReportError("No existing Pass/PASS Result cell was found in the current testcase block")
        if status_style_row is not None:
            copy_cell_style(ws.cell(status_style_row, result_column), ws.cell(row, result_column))

        definition_values: dict[str, Any] = {
            "tc_id": testcase["tc_id"],
            "module": testcase["module"],
            "test_title": testcase["test_title"],
            "description": testcase["description"],
            "test_steps": format_steps(testcase["test_steps"]),
            "test_data": testcase["test_data"],
            "expected_result": testcase["expected_result"],
        }
        execution_values: dict[str, Any] = {
            "actual_result": testcase["actual_result"],
            "evidence": evidence_text(testcase["evidence"]),
            "status": testcase["status"],
        }
        values_to_write = execution_values if existing_row is not None else {
            **definition_values, **execution_values
        }
        preserved_definitions = tuple(
            (FIELD_TO_HEADER[field], ws.cell(row, headers[normalized(FIELD_TO_HEADER[field])]).value)
            for field in definition_fields
        )
        feedback_values = tuple(
            (feedback_header, ws.cell(row, headers[normalized(feedback_header)]).value)
            for feedback_header in ("Phản hồi", "FEEDBACK")
            if normalized(feedback_header) in headers
        )
        prior_screenshot_count = anchored_image_count(ws, row, evidence_column)
        for field, value in values_to_write.items():
            target = ws.cell(row, headers[normalized(FIELD_TO_HEADER[field])])
            if is_formula(target.value):
                raise ReportError(f"Refusing to overwrite formula in {ws.title}!{target.coordinate}")
            target.value = value

        text_line_count = max(1, str(execution_values["evidence"]).count("\n") + 1)
        screenshot_count = add_screenshots(
            ws, row, evidence_column, testcase["evidence"], text_line_count
        )
        non_image_count = sum(1 for item in testcase["evidence"] if item["type"] != "screenshot")
        handling = f"{non_image_count} text evidence; {screenshot_count} screenshot image(s)"
        records.append(
            WriteRecord(
                row=row,
                tc_id=testcase["tc_id"],
                status=testcase["status"],
                expected_result=testcase["expected_result"],
                actual_result=testcase["actual_result"],
                evidence_text=execution_values["evidence"],
                evidence_handling=handling,
                screenshot_count=screenshot_count,
                prior_screenshot_count=prior_screenshot_count,
                operation=operation,
                definition_values=(
                    preserved_definitions if existing_row is not None else
                    tuple((FIELD_TO_HEADER[field], definition_values[field]) for field in definition_fields)
                ),
                feedback_values=feedback_values,
            )
        )
        existing_ids[testcase["tc_id"]] = row
    return records, header_row, headers


def write_bugs(
    workbook: object, validated: ValidatedInput
) -> tuple[list[BugWriteRecord], int, dict[str, int]]:
    if BUG_SHEET not in workbook.sheetnames:
        raise ReportError("Required sheet not found: NoteBug")
    ws = workbook[BUG_SHEET]
    header_row, bug_id_column = find_last_header(ws, "Bug ID")
    headers = header_columns(ws, header_row, BUG_FIELD_TO_HEADER)
    mapped_columns = sorted(
        {headers[normalized(header)] for header in BUG_FIELD_TO_HEADER.values()}
    )
    bugs: list[dict[str, Any]] = validated.document["bugs"]
    existing_ids = unique_identifier_rows(ws, header_row, bug_id_column, "Bug ID")
    environment: dict[str, str] = validated.document["run"].get("environment", {})
    evidence_column = headers[normalized("Log/Screenshot/Video")]
    records: list[BugWriteRecord] = []
    for bug in bugs:
        existing_row = existing_ids.get(bug["bug_id"])
        operation = "updated existing" if existing_row is not None else "appended new"
        if existing_row is not None:
            row = existing_row
        else:
            row, last_bug_row = next_append_row(
                ws, header_row, bug_id_column, "Bug ID", mapped_columns
            )
            if not row_is_blank(ws, row, mapped_columns):
                raise ReportError(f"Target row is not blank in mapped columns: NoteBug!{row}")
            source_row = nearest_style_row(
                ws, header_row, row, last_bug_row, mapped_columns
            )
            prepare_row_style_if_unstyled(ws, source_row, row, mapped_columns)

        values: dict[str, Any] = {
            "bug_id": bug["bug_id"],
            "reporter": environment["reporter"],
            "device": environment["device"],
            "os_app_version": environment["os_app_version"],
            "module": bug["module"],
            "linked_tc_id": bug["linked_tc_id"],
            "title": bug["title"],
            "reproduction_steps": format_steps(bug["reproduction_steps"]),
            "actual_result": bug["actual_result"],
            "expected_result": bug["expected_result"],
            "evidence": evidence_text(bug["evidence"]),
        }
        prior_screenshot_count = anchored_image_count(ws, row, evidence_column)
        for field, value in values.items():
            target = ws.cell(row, headers[normalized(BUG_FIELD_TO_HEADER[field])])
            if is_formula(target.value):
                raise ReportError(f"Refusing to overwrite formula in {ws.title}!{target.coordinate}")
            target.value = value

        text_line_count = max(1, str(values["evidence"]).count("\n") + 1)
        screenshot_count = add_screenshots(
            ws, row, evidence_column, bug["evidence"], text_line_count
        )
        non_image_count = sum(1 for item in bug["evidence"] if item["type"] != "screenshot")
        handling = f"{non_image_count} text evidence; {screenshot_count} screenshot image(s)"
        records.append(
            BugWriteRecord(
                row=row,
                bug_id=bug["bug_id"],
                linked_tc_id=bug["linked_tc_id"],
                actual_result=bug["actual_result"],
                expected_result=bug["expected_result"],
                reporter=environment["reporter"],
                device=environment["device"],
                os_app_version=environment["os_app_version"],
                module=bug["module"],
                title=bug["title"],
                reproduction_steps=values["reproduction_steps"],
                evidence_text=values["evidence"],
                evidence_handling=handling,
                screenshot_count=screenshot_count,
                prior_screenshot_count=prior_screenshot_count,
                operation=operation,
            )
        )
        existing_ids[bug["bug_id"]] = row
    return records, header_row, headers


def compute_current_metrics(workbook: object) -> SummaryMetrics:
    if TESTCASE_SHEET not in workbook.sheetnames or BUG_SHEET not in workbook.sheetnames:
        raise ReportError("Cannot compute Summary: Testcase or NoteBug sheet is missing")
    testcase_ws = workbook[TESTCASE_SHEET]
    testcase_header_row, tc_id_column = find_last_header(testcase_ws, "TC_ID")
    testcase_headers = header_columns(testcase_ws, testcase_header_row, FIELD_TO_HEADER)
    result_column = testcase_headers[normalized("Result")]
    testcase_rows = current_identifier_rows(
        testcase_ws, testcase_header_row, tc_id_column, "TC_ID"
    )
    passed = 0
    failed = 0
    unverified = 0
    for row in testcase_rows:
        value = testcase_ws.cell(row, result_column).value
        if value in (None, ""):
            unverified += 1
            continue
        status = str(value).strip().upper()
        if status == "PASS":
            passed += 1
        elif status == "FAIL":
            failed += 1
        elif status == "UNVERIFIED":
            unverified += 1
        else:
            raise ReportError(
                f"Unsupported Result in current Testcase block at "
                f"{testcase_ws.title}!{testcase_ws.cell(row, result_column).coordinate}: {value!r}"
            )

    bug_ws = workbook[BUG_SHEET]
    bug_header_row, bug_id_column = find_last_header(bug_ws, "Bug ID")
    bug_rows = current_identifier_rows(bug_ws, bug_header_row, bug_id_column, "Bug ID")
    total = len(testcase_rows)
    return SummaryMetrics(
        total=total,
        passed=passed,
        failed=failed,
        unverified=unverified,
        bugs=len(bug_rows),
        pass_rate=(passed / total if total else 0.0),
    )


def summary_label_metric(value: object) -> str | None:
    label = normalized(value).rstrip(":").strip()
    for metric, aliases in SUMMARY_LABELS.items():
        if label in aliases:
            return metric
    return None


def summary_value_cell(ws: Worksheet, label_cell: object) -> object:
    value_column = label_cell.column + 1
    for merged_range in ws.merged_cells.ranges:
        if label_cell.coordinate in merged_range:
            value_column = merged_range.max_col + 1
            break
    return ws.cell(label_cell.row, value_column)


def find_current_summary_kpis(ws: Worksheet) -> dict[str, object]:
    candidates: list[tuple[int, int, str, object]] = []
    for row in ws.iter_rows():
        for cell in row:
            metric = summary_label_metric(cell.value)
            if metric is not None:
                candidates.append((cell.row, cell.column, metric, cell))
    if not candidates:
        raise ReportError("No Summary KPI labels were found")

    clusters: list[list[tuple[int, int, str, object]]] = []
    for candidate in sorted(candidates):
        starts_repeated_block = (
            bool(clusters)
            and candidate[2] == "total"
            and any(item[2] == "total" for item in clusters[-1])
        )
        if (
            not clusters
            or candidate[0] - clusters[-1][-1][0] > 10
            or starts_repeated_block
        ):
            clusters.append([candidate])
        else:
            clusters[-1].append(candidate)
    eligible = [
        cluster for cluster in clusters
        if set(REQUIRED_SUMMARY_KPIS).issubset({item[2] for item in cluster})
    ]
    if not eligible:
        found = sorted({item[2] for item in candidates})
        raise ReportError(
            "Current Summary block is missing required KPI label(s); "
            f"found metrics: {found}"
        )
    current = eligible[-1]
    kpis: dict[str, object] = {}
    for _, _, metric, label_cell in current:
        if metric in kpis:
            raise ReportError(
                f"Duplicate Summary KPI label for {metric!r} in current block"
            )
        kpis[metric] = summary_value_cell(ws, label_cell)
    return kpis


def write_summary(workbook: object) -> SummaryWriteRecord:
    if SUMMARY_SHEET not in workbook.sheetnames:
        raise ReportError("Required sheet not found: Summary")
    metrics = compute_current_metrics(workbook)
    kpis = find_current_summary_kpis(workbook[SUMMARY_SHEET])
    values: dict[str, int | float] = {
        "total": metrics.total,
        "pass": metrics.passed,
        "fail": metrics.failed,
        "unverified": metrics.unverified,
        "bugs": metrics.bugs,
        "pass_rate": metrics.pass_rate,
    }
    formula_replacements: set[str] = set()
    for metric, cell in kpis.items():
        if is_formula(cell.value):
            formula_replacements.add(cell.coordinate)
        cell.value = values[metric]
    return SummaryWriteRecord(
        metrics=metrics,
        cells={metric: cell.coordinate for metric, cell in kpis.items()},
        formula_replacements=frozenset(formula_replacements),
    )


def unexpected_formula_differences(
    before: WorkbookSnapshot,
    after: WorkbookSnapshot,
    allowed_summary_cells: set[str],
) -> list[str]:
    differences: list[str] = []
    for sheet_name in sorted(set(before.formula_states) | set(after.formula_states)):
        old = dict(before.formula_states.get(sheet_name, ()))
        new = dict(after.formula_states.get(sheet_name, ()))
        for coordinate in sorted(set(old) | set(new)):
            if old.get(coordinate) == new.get(coordinate):
                continue
            if sheet_name == SUMMARY_SHEET and coordinate in allowed_summary_cells:
                continue
            differences.append(
                f"{sheet_name}!{coordinate}: before={old.get(coordinate)!r}, "
                f"after={new.get(coordinate)!r}"
            )
    return differences


def compare_equal_counts(
    label: str, before: Mapping[str, int], after: Mapping[str, int], failures: list[str]
) -> None:
    if dict(before) != dict(after):
        failures.append(f"{label} changed: before={dict(before)}, after={dict(after)}")


def verify_output(
    workbook: object,
    before: WorkbookSnapshot,
    testcase_records: Sequence[WriteRecord],
    testcase_header_row: int,
    testcase_headers: Mapping[str, int],
    bug_records: Sequence[BugWriteRecord],
    bug_header_row: int,
    bug_headers: Mapping[str, int],
    summary_record: SummaryWriteRecord,
) -> tuple[WorkbookSnapshot, list[str], bool, list[str]]:
    after = snapshot(workbook)
    failures: list[str] = []
    if after.sheet_names != before.sheet_names:
        failures.append(f"sheet names changed: before={before.sheet_names}, after={after.sheet_names}")
    inserted_images = {
        TESTCASE_SHEET: sum(record.screenshot_count for record in testcase_records),
        BUG_SHEET: sum(record.screenshot_count for record in bug_records),
    }
    for sheet_name, count_before in before.image_counts.items():
        count_after = after.image_counts.get(sheet_name, -1)
        expected_after = count_before + inserted_images.get(sheet_name, 0)
        if count_after != expected_after:
            failures.append(
                f"image count mismatch on {sheet_name}: before={count_before}, "
                f"inserted={inserted_images.get(sheet_name, 0)}, "
                f"expected={expected_after}, after={count_after}"
            )
        old_signatures = list(before.image_signatures.get(sheet_name, ()))
        new_signatures = list(after.image_signatures.get(sheet_name, ()))
        for signature in old_signatures:
            if signature not in new_signatures:
                failures.append(
                    f"existing image anchor/size disappeared on {sheet_name}: {signature}"
                )
            else:
                new_signatures.remove(signature)
    compare_equal_counts("merged-range counts", before.merged_counts, after.merged_counts, failures)
    if before.merged_ranges != after.merged_ranges:
        failures.append("merged ranges changed")
    unexpected_formulas = unexpected_formula_differences(
        before, after, set(summary_record.formula_replacements)
    )
    failures.extend(f"unexpected formula change: {change}" for change in unexpected_formulas)
    summary_differences = summary_semantic_differences(
        before.summary_state, after.summary_state, set(summary_record.cells.values())
    )
    summary_preserved = not summary_differences
    failures.extend(summary_differences)

    if TESTCASE_SHEET not in workbook.sheetnames:
        failures.append("required sheet is missing: Testcase")
        return after, failures, summary_preserved, unexpected_formulas
    ws = workbook[TESTCASE_SHEET]
    tc_id_column = testcase_headers[normalized("TC_ID")]
    result_column = testcase_headers[normalized("Result")]
    expected_column = testcase_headers[normalized("Expected Result")]
    actual_column = testcase_headers[normalized("Actual Result")]
    testcase_evidence_column = testcase_headers[normalized("Log/Screenshot/Video")]
    for record in testcase_records:
        occurrences = sum(
            str(ws.cell(row, tc_id_column).value) == record.tc_id
            for row in range(testcase_header_row + 1, ws.max_row + 1)
            if ws.cell(row, tc_id_column).value not in (None, "")
        )
        if occurrences != 1:
            failures.append(f"{record.tc_id} occurs {occurrences} time(s) in the current block")
        checks = (
            ("Actual Result", ws.cell(record.row, actual_column).value, record.actual_result),
            ("Result", ws.cell(record.row, result_column).value, record.status),
            ("Expected Result", ws.cell(record.row, expected_column).value, record.expected_result),
            (
                "Log/Screenshot/Video",
                ws.cell(record.row, testcase_evidence_column).value,
                record.evidence_text,
            ),
        )
        for label, actual, expected in checks:
            if not verification_values_equal(actual, expected):
                failures.append(
                    f"{record.tc_id} {label} mismatch: expected={expected!r}, actual={actual!r}"
                )
        for header, expected in record.definition_values:
            column = testcase_headers[normalized(header)]
            actual = ws.cell(record.row, column).value
            if not verification_values_equal(actual, expected):
                failures.append(
                    f"{record.tc_id} preserved {header} mismatch: "
                    f"expected={expected!r}, actual={actual!r}"
                )
        anchored_screenshots = sum(
            image_anchor_signature(image)[:2] == (record.row - 1, testcase_evidence_column - 1)
            for image in getattr(ws, "_images", ())
        )
        expected_screenshots = record.prior_screenshot_count + record.screenshot_count
        if anchored_screenshots != expected_screenshots:
            failures.append(
                f"{record.tc_id} screenshot anchor count mismatch: "
                f"expected={expected_screenshots}, actual={anchored_screenshots}"
            )
        for feedback_header, expected in record.feedback_values:
            feedback_column = testcase_headers[normalized(feedback_header)]
            actual = ws.cell(record.row, feedback_column).value
            if not verification_values_equal(actual, expected):
                failures.append(
                    f"{record.tc_id} {feedback_header} changed: "
                    f"expected={expected!r}, actual={actual!r}"
                )

    if BUG_SHEET not in workbook.sheetnames:
        failures.append("required sheet is missing: NoteBug")
        return after, failures, summary_preserved, unexpected_formulas
    bug_ws = workbook[BUG_SHEET]
    bug_id_column = bug_headers[normalized("Bug ID")]
    linked_tc_column = bug_headers[normalized("Testcase liên quan")]
    bug_actual_column = bug_headers[normalized("Kết quả thực tế")]
    bug_expected_column = bug_headers[normalized("Kết quả mong muốn")]
    bug_module_column = bug_headers[normalized("Module")]
    bug_title_column = bug_headers[normalized("Tiêu đề bug")]
    bug_steps_column = bug_headers[normalized("Bước tái hiện")]
    bug_evidence_column = bug_headers[normalized("Log/Screenshot/Video")]
    environment_columns = {
        "Người report": bug_headers[normalized("Người report")],
        "Thiết bị": bug_headers[normalized("Thiết bị")],
        "OS/App version": bug_headers[normalized("OS/App version")],
    }
    for record in bug_records:
        occurrences = sum(
            str(bug_ws.cell(row, bug_id_column).value) == record.bug_id
            for row in range(bug_header_row + 1, bug_ws.max_row + 1)
            if bug_ws.cell(row, bug_id_column).value not in (None, "")
        )
        if occurrences != 1:
            failures.append(
                f"{record.bug_id} occurs {occurrences} time(s) in the current NoteBug block"
            )
        checks = (
            ("Testcase liên quan", bug_ws.cell(record.row, linked_tc_column).value, record.linked_tc_id),
            ("Kết quả thực tế", bug_ws.cell(record.row, bug_actual_column).value, record.actual_result),
            ("Kết quả mong muốn", bug_ws.cell(record.row, bug_expected_column).value, record.expected_result),
            ("Người report", bug_ws.cell(record.row, environment_columns["Người report"]).value, record.reporter),
            ("Thiết bị", bug_ws.cell(record.row, environment_columns["Thiết bị"]).value, record.device),
            ("OS/App version", bug_ws.cell(record.row, environment_columns["OS/App version"]).value, record.os_app_version),
            ("Module", bug_ws.cell(record.row, bug_module_column).value, record.module),
            ("Tiêu đề bug", bug_ws.cell(record.row, bug_title_column).value, record.title),
            ("Bước tái hiện", bug_ws.cell(record.row, bug_steps_column).value, record.reproduction_steps),
            ("Log/Screenshot/Video", bug_ws.cell(record.row, bug_evidence_column).value, record.evidence_text),
        )
        for label, actual, expected in checks:
            if not verification_values_equal(actual, expected):
                failures.append(
                    f"{record.bug_id} {label} mismatch: expected={expected!r}, actual={actual!r}"
                )
        anchored_screenshots = sum(
            image_anchor_signature(image)[:2] == (record.row - 1, bug_evidence_column - 1)
            for image in getattr(bug_ws, "_images", ())
        )
        expected_screenshots = record.prior_screenshot_count + record.screenshot_count
        if anchored_screenshots != expected_screenshots:
            failures.append(
                f"{record.bug_id} screenshot anchor count mismatch: "
                f"expected={expected_screenshots}, actual={anchored_screenshots}"
            )

    try:
        if SUMMARY_SHEET not in workbook.sheetnames:
            raise ReportError("Required sheet is missing: Summary")
        recomputed = compute_current_metrics(workbook)
        if recomputed != summary_record.metrics:
            failures.append(
                f"Summary recomputed metrics changed: "
                f"expected={summary_record.metrics!r}, actual={recomputed!r}"
            )
        summary_kpis = find_current_summary_kpis(workbook[SUMMARY_SHEET])
        expected_values: dict[str, int | float] = {
            "total": recomputed.total,
            "pass": recomputed.passed,
            "fail": recomputed.failed,
            "unverified": recomputed.unverified,
            "bugs": recomputed.bugs,
            "pass_rate": recomputed.pass_rate,
        }
        for metric, cell in summary_kpis.items():
            expected = expected_values[metric]
            actual = cell.value
            matches = (
                isinstance(expected, float)
                and isinstance(actual, (int, float))
                and not isinstance(actual, bool)
                and math.isclose(float(actual), expected, rel_tol=1e-12, abs_tol=1e-12)
            ) or verification_values_equal(actual, expected)
            if not matches:
                failures.append(
                    f"Summary KPI {metric} mismatch at {cell.coordinate}: "
                    f"expected={expected!r}, actual={actual!r}"
                )
    except ReportError as exc:
        failures.append(str(exc))
    return after, failures, summary_preserved, unexpected_formulas


def print_counts(label: str, before: Mapping[str, int], after: Mapping[str, int]) -> None:
    print(f"- {label}:")
    for sheet_name in before:
        print(f"  {sheet_name}: {before[sheet_name]} / {after.get(sheet_name, 'missing')}")


def existing_images_are_preserved(before: WorkbookSnapshot, after: WorkbookSnapshot) -> bool:
    for sheet_name, old_signatures in before.image_signatures.items():
        remaining = list(after.image_signatures.get(sheet_name, ()))
        for signature in old_signatures:
            if signature not in remaining:
                return False
            remaining.remove(signature)
    return True


def print_report(
    input_path: Path,
    template: Path,
    output: Path,
    testcase_records: Sequence[WriteRecord],
    bug_records: Sequence[BugWriteRecord],
    before: WorkbookSnapshot,
    after: WorkbookSnapshot,
    summary_record: SummaryWriteRecord,
    summary_preserved: bool,
    unexpected_formulas: Sequence[str],
    failures: Sequence[str],
) -> None:
    print("REPORT WRITE")
    print(f"- input JSON: {input_path}")
    print(f"- template: {template}")
    print(f"- output: {output}")
    print("- testcase:")
    print(
        "  updated existing: "
        f"{sum(record.operation == 'updated existing' for record in testcase_records)}"
    )
    print(
        "  appended new: "
        f"{sum(record.operation == 'appended new' for record in testcase_records)}"
    )
    print("- bugs:")
    print(
        "  updated existing: "
        f"{sum(record.operation == 'updated existing' for record in bug_records)}"
    )
    print(
        "  appended new: "
        f"{sum(record.operation == 'appended new' for record in bug_records)}"
    )
    screenshot_total = sum(record.screenshot_count for record in testcase_records)
    screenshot_total += sum(record.screenshot_count for record in bug_records)
    print(f"- screenshots inserted: {screenshot_total}")
    metrics = summary_record.metrics
    print("- summary:")
    print(f"  total: {metrics.total}")
    print(f"  pass: {metrics.passed}")
    print(f"  fail: {metrics.failed}")
    print(f"  unverified: {metrics.unverified}")
    print(f"  bugs: {metrics.bugs}")
    print(f"  pass rate: {metrics.pass_rate:.6%}")
    print("- template unchanged: YES")
    print(
        "- existing images preserved: "
        f"{'YES' if existing_images_are_preserved(before, after) else 'NO'}"
    )
    merged_preserved = (
        before.merged_counts == after.merged_counts
        and before.merged_ranges == after.merged_ranges
    )
    print(f"- merged ranges preserved: {'YES' if merged_preserved else 'NO'}")
    print(
        "- Summary non-KPI semantic state preserved: "
        f"{'YES' if summary_preserved else 'NO'}"
    )
    print("- unexpected formula changes:")
    if unexpected_formulas:
        for change in unexpected_formulas:
            print(f"  {change}")
    else:
        print("  none")
    print(f"- verification: {'FAIL' if failures else 'PASS'}")
    for failure in failures:
        print(f"  ERROR: {failure}")


def run_write(input_path: Path, template_path: Path, output_path: Path) -> bool:
    # Contract validation deliberately precedes all copy/write operations.
    validated = load_and_validate_json(input_path)
    template, output = validate_paths(template_path, output_path)
    output.parent.mkdir(parents=True, exist_ok=True)

    template_hash = file_sha256(template)
    shutil.copy2(template, output)
    if file_sha256(output) != template_hash:
        raise ReportError("Copied output does not match the source template")

    workbook = load_workbook(output, read_only=False, data_only=False)
    try:
        before = snapshot(workbook)
        testcase_records, testcase_header_row, testcase_headers = write_testcases(
            workbook, validated
        )
        bug_records, bug_header_row, bug_headers = write_bugs(workbook, validated)
        summary_record = write_summary(workbook)
        workbook.save(output)
    finally:
        workbook.close()

    if file_sha256(template) != template_hash:
        raise ReportError("Template content changed during report generation")

    reopened = load_workbook(output, read_only=False, data_only=False)
    try:
        after, failures, summary_preserved, unexpected_formulas = verify_output(
            reopened,
            before,
            testcase_records,
            testcase_header_row,
            testcase_headers,
            bug_records,
            bug_header_row,
            bug_headers,
            summary_record,
        )
    finally:
        reopened.close()

    print_report(
        validated.path,
        template,
        output,
        testcase_records,
        bug_records,
        before,
        after,
        summary_record,
        summary_preserved,
        unexpected_formulas,
        failures,
    )
    return not failures


def parse_args(argv: Sequence[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Validate qa-result.json, upsert current Testcase/NoteBug blocks, "
            "and recompute label-based Summary KPIs in a copied QA workbook."
        )
    )
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--template", type=Path, default=DEFAULT_TEMPLATE)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    return parser.parse_args(argv)


def main(argv: Sequence[str] | None = None) -> int:
    args = parse_args(argv)
    try:
        return 0 if run_write(args.input, args.template, args.output) else 1
    except ReportError as exc:
        print("REPORT WRITE", file=sys.stderr)
        print("- verification: FAIL", file=sys.stderr)
        print(f"- error: {exc}", file=sys.stderr)
        return 2
    except Exception as exc:
        print("REPORT WRITE", file=sys.stderr)
        print("- verification: FAIL", file=sys.stderr)
        print(f"- error: unexpected failure: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
