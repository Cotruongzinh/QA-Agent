"""Safely write one smoke-test testcase to a copied QA report workbook.

This script never edits the source template. It validates the paths, copies the
template to the output with shutil.copy2, and loads/saves only the output copy.
"""

from __future__ import annotations

import argparse
import copy
import hashlib
import json
import shutil
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Mapping, Sequence

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet


DEFAULT_TEMPLATE = Path(r"C:\QA-Agent\templates\QA_Report_Template.xlsx")
DEFAULT_OUTPUT = Path(r"C:\QA-Agent\output\QA_Report_SMOKE.xlsx")
DEFAULT_TEMPLATES_DIR = Path(r"C:\QA-Agent\templates")

SMOKE_ID = "TC_SMOKE_001"
REQUIRED_SHEETS = ("Testcase", "NoteBug", "Summary")
EVIDENCE_HEADER = "Log/Screenshot/Video"

SMOKE_VALUES: dict[str, str] = {
    "TC_ID": SMOKE_ID,
    "Module": "SMOKE",
    "Test Title": "Excel report write smoke test",
    "Description/Test Summary": (
        "Kiểm tra khả năng ghi báo cáo trên bản copy mà không thay đổi template gốc."
    ),
    "Test Steps": (
        "1. Copy template sang output.\n"
        "2. Ghi một testcase smoke.\n"
        "3. Save và mở lại output để xác minh."
    ),
    "Test Data": "N/A",
    "Expected Result": (
        "Dữ liệu smoke được ghi vào file output và template gốc không thay đổi."
    ),
    "Actual Result": "Smoke write successful.",
    "Result": "Pass",
    "Phản hồi": "SMOKE TEST - không phải kết quả kiểm thử thật.",
    "FEEDBACK": "Tạo tự động bởi excel-report smoke test.",
}


class SmokeTestError(RuntimeError):
    """Raised when a safety rule or smoke-test verification fails."""


@dataclass(frozen=True)
class WorkbookSnapshot:
    sheet_names: tuple[str, ...]
    image_counts: dict[str, int]
    merged_counts: dict[str, int]
    formula_counts: dict[str, int]
    protected_values: dict[str, dict[str, object]]


def normalized(value: object) -> str:
    if value is None:
        return ""
    return " ".join(str(value).strip().casefold().split())


def resolved(path: Path) -> Path:
    return path.expanduser().resolve(strict=False)


def is_within(path: Path, directory: Path) -> bool:
    try:
        path.relative_to(directory)
        return True
    except ValueError:
        return False


def is_in_templates_directory(path: Path) -> bool:
    if is_within(path, resolved(DEFAULT_TEMPLATES_DIR)):
        return True
    return any(parent.name.casefold() == "templates" for parent in path.parents)


def validate_paths(template_path: Path, output_path: Path) -> tuple[Path, Path]:
    template = resolved(template_path)
    output = resolved(output_path)
    if template == output:
        raise SmokeTestError("Refusing to run: input path equals output path")
    if is_in_templates_directory(output):
        raise SmokeTestError("Refusing to run: output is inside a templates directory")
    if not template.is_file():
        raise SmokeTestError(f"Template not found: {template}")
    if template.suffix.casefold() != ".xlsx" or output.suffix.casefold() != ".xlsx":
        raise SmokeTestError("Input and output must both be .xlsx files")
    return template, output


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def is_formula(value: object) -> bool:
    return isinstance(value, str) and value.startswith("=")


def formula_count(ws: Worksheet) -> int:
    return sum(
        is_formula(cell.value)
        for row in ws.iter_rows()
        for cell in row
        if cell.value is not None
    )


def cell_values(ws: Worksheet) -> dict[str, object]:
    return {
        cell.coordinate: cell.value
        for row in ws.iter_rows()
        for cell in row
        if cell.value is not None
    }


def snapshot(workbook: object) -> WorkbookSnapshot:
    return WorkbookSnapshot(
        sheet_names=tuple(workbook.sheetnames),
        image_counts={
            ws.title: len(getattr(ws, "_images", ())) for ws in workbook.worksheets
        },
        merged_counts={
            ws.title: len(ws.merged_cells.ranges) for ws in workbook.worksheets
        },
        formula_counts={ws.title: formula_count(ws) for ws in workbook.worksheets},
        protected_values={
            sheet_name: cell_values(workbook[sheet_name])
            for sheet_name in ("NoteBug", "Summary")
            if sheet_name in workbook.sheetnames
        },
    )


def tc_id_occurrences(ws: Worksheet) -> list[tuple[int, int]]:
    return [
        (cell.row, cell.column)
        for row in ws.iter_rows()
        for cell in row
        if normalized(cell.value) == normalized("TC_ID")
    ]


def header_columns(ws: Worksheet, header_row: int) -> dict[str, int]:
    columns: dict[str, int] = {}
    duplicates: set[str] = set()
    for column in range(1, ws.max_column + 1):
        label = normalized(ws.cell(header_row, column).value)
        if not label:
            continue
        if label in columns:
            duplicates.add(label)
        else:
            columns[label] = column
    required = {normalized(name): name for name in SMOKE_VALUES}
    missing = [display for key, display in required.items() if key not in columns]
    ambiguous = [required[key] for key in duplicates if key in required]
    if missing:
        raise SmokeTestError("Missing testcase header(s): " + ", ".join(missing))
    if ambiguous:
        raise SmokeTestError("Duplicate testcase header(s): " + ", ".join(ambiguous))
    return columns


def find_target_row(
    ws: Worksheet, header_row: int, tc_id_column: int
) -> tuple[int, int | None]:
    testcase_rows = [
        row
        for row in range(header_row + 1, ws.max_row + 1)
        if ws.cell(row, tc_id_column).value not in (None, "")
        and normalized(ws.cell(row, tc_id_column).value) != normalized("TC_ID")
    ]
    last_testcase_row = max(testcase_rows) if testcase_rows else None
    return (
        last_testcase_row + 1 if last_testcase_row is not None else header_row + 1,
        last_testcase_row,
    )


STYLE_ATTRIBUTES = (
    "font",
    "fill",
    "border",
    "alignment",
    "number_format",
    "protection",
)


def copy_cell_style(source: object, target: object) -> None:
    for attribute in STYLE_ATTRIBUTES:
        setattr(target, attribute, copy.copy(getattr(source, attribute)))


def prepare_target_style(
    ws: Worksheet,
    target_row: int,
    previous_testcase_row: int | None,
    header_map: Mapping[str, int],
    header_row: int,
) -> None:
    style_width = max(ws.max_column, max(header_map.values()))
    target_has_style = any(
        ws.cell(target_row, column).has_style for column in range(1, style_width + 1)
    )
    if target_has_style:
        return
    if previous_testcase_row is None:
        raise SmokeTestError(
            "Target row is unstyled and there is no previous testcase row to copy"
        )

    for column in range(1, style_width + 1):
        copy_cell_style(
            ws.cell(previous_testcase_row, column), ws.cell(target_row, column)
        )
    previous_height = ws.row_dimensions[previous_testcase_row].height
    if previous_height is not None:
        ws.row_dimensions[target_row].height = previous_height

    # Prefer an existing Pass cell as the Result style exemplar when available.
    result_column = header_map[normalized("Result")]
    tc_id_column = header_map[normalized("TC_ID")]
    pass_rows = [
        row
        for row in range(header_row + 1, target_row)
        if normalized(ws.cell(row, tc_id_column).value) != ""
        and normalized(ws.cell(row, result_column).value) == normalized("Pass")
    ]
    if pass_rows:
        copy_cell_style(
            ws.cell(pass_rows[-1], result_column), ws.cell(target_row, result_column)
        )


def ensure_writable_target(
    ws: Worksheet, target_row: int, header_map: Mapping[str, int]
) -> None:
    for header in SMOKE_VALUES:
        cell = ws.cell(target_row, header_map[normalized(header)])
        if is_formula(cell.value):
            raise SmokeTestError(f"Refusing to overwrite formula in {ws.title}!{cell.coordinate}")
        if cell.value not in (None, ""):
            raise SmokeTestError(f"Target cell is not blank: {ws.title}!{cell.coordinate}")


def count_smoke_ids(workbook: object) -> int:
    return sum(
        normalized(cell.value) == normalized(SMOKE_ID)
        for ws in workbook.worksheets
        for row in ws.iter_rows()
        for cell in row
    )


def compare_counts(
    label: str, before: Mapping[str, int], after: Mapping[str, int], failures: list[str]
) -> None:
    if dict(before) != dict(after):
        failures.append(f"{label} changed: before={dict(before)}, after={dict(after)}")


def verify_output(
    workbook: object,
    before: WorkbookSnapshot,
    target_row: int,
    header_map: Mapping[str, int],
) -> tuple[WorkbookSnapshot, list[str]]:
    after = snapshot(workbook)
    failures: list[str] = []
    if after.sheet_names != before.sheet_names:
        failures.append(
            f"sheet names changed: before={before.sheet_names}, after={after.sheet_names}"
        )
    compare_counts("embedded image counts", before.image_counts, after.image_counts, failures)
    compare_counts("merged-range counts", before.merged_counts, after.merged_counts, failures)
    compare_counts("formula counts", before.formula_counts, after.formula_counts, failures)

    missing_sheets = [name for name in REQUIRED_SHEETS if name not in workbook.sheetnames]
    if missing_sheets:
        failures.append("required sheet(s) missing: " + ", ".join(missing_sheets))
    if count_smoke_ids(workbook) != 1:
        failures.append(f"{SMOKE_ID} does not exist exactly once")

    if "Testcase" in workbook.sheetnames:
        ws = workbook["Testcase"]
        result = ws.cell(target_row, header_map[normalized("Result")]).value
        if normalized(result) != normalized("Pass"):
            failures.append(f"Result is not Pass at Testcase row {target_row}: {result!r}")
        for header, expected in SMOKE_VALUES.items():
            actual = ws.cell(target_row, header_map[normalized(header)]).value
            if actual != expected:
                failures.append(
                    f"value mismatch for {header!r}: expected={expected!r}, actual={actual!r}"
                )

    for sheet_name, values_before in before.protected_values.items():
        if sheet_name in workbook.sheetnames:
            values_after = cell_values(workbook[sheet_name])
            if values_after != values_before:
                failures.append(f"protected sheet values changed: {sheet_name}")
    return after, failures


def print_count_report(
    label: str, before: Mapping[str, int], after: Mapping[str, int]
) -> None:
    print(f"- {label}:")
    for sheet_name in before:
        print(
            f"  {sheet_name}: {before[sheet_name]} / "
            f"{after.get(sheet_name, 'missing')}"
        )


def print_report(
    template: Path,
    output: Path,
    target_row: int,
    before: WorkbookSnapshot,
    after: WorkbookSnapshot,
    failures: Sequence[str],
) -> None:
    print("WRITE SMOKE TEST")
    print(f"- template path: {template}")
    print(f"- output path: {output}")
    print(f"- target row: {target_row}")
    print("- values written:")
    for header, value in SMOKE_VALUES.items():
        print(f"  {header}: {json.dumps(value, ensure_ascii=False)}")
    print(f"  {EVIDENCE_HEADER}: (not written)")
    print_count_report("images before/after per sheet", before.image_counts, after.image_counts)
    print_count_report("merged ranges before/after", before.merged_counts, after.merged_counts)
    print_count_report("formulas before/after", before.formula_counts, after.formula_counts)
    print(f"- verification: {'FAIL' if failures else 'PASS'}")
    for failure in failures:
        print(f"  ERROR: {failure}")


def run_smoke_test(template_path: Path, output_path: Path) -> bool:
    template, output = validate_paths(template_path, output_path)
    output.parent.mkdir(parents=True, exist_ok=True)

    # This is deliberately the first workbook operation. Only the copied output
    # is ever passed to openpyxl.
    shutil.copy2(template, output)
    template_hash = file_sha256(template)
    if file_sha256(output) != template_hash:
        raise SmokeTestError("Copied output does not match the source template")

    workbook = load_workbook(output, read_only=False, data_only=False)
    target_row = 0
    header_map: dict[str, int] = {}
    try:
        if "Testcase" not in workbook.sheetnames:
            raise SmokeTestError("Required sheet not found: Testcase")
        before = snapshot(workbook)
        ws = workbook["Testcase"]
        occurrences = tc_id_occurrences(ws)
        if not occurrences:
            raise SmokeTestError("TC_ID header not found on sheet Testcase")
        header_row, tc_id_column = occurrences[-1]
        header_map = header_columns(ws, header_row)

        if count_smoke_ids(workbook):
            raise SmokeTestError(f"Refusing to duplicate existing {SMOKE_ID}")
        target_row, previous_testcase_row = find_target_row(
            ws, header_row, tc_id_column
        )
        ensure_writable_target(ws, target_row, header_map)
        prepare_target_style(
            ws, target_row, previous_testcase_row, header_map, header_row
        )
        for header, value in SMOKE_VALUES.items():
            ws.cell(target_row, header_map[normalized(header)]).value = value

        workbook.save(output)
    finally:
        workbook.close()

    if file_sha256(template) != template_hash:
        raise SmokeTestError("Template content changed during the smoke test")

    reopened = load_workbook(output, read_only=False, data_only=False)
    try:
        after, failures = verify_output(reopened, before, target_row, header_map)
    finally:
        reopened.close()

    print_report(template, output, target_row, before, after, failures)
    return not failures


def parse_args(argv: Sequence[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Write and verify one safe testcase in a copied QA report."
    )
    parser.add_argument("--template", type=Path, default=DEFAULT_TEMPLATE)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    return parser.parse_args(argv)


def main(argv: Sequence[str] | None = None) -> int:
    args = parse_args(argv)
    try:
        return 0 if run_smoke_test(args.template, args.output) else 1
    except SmokeTestError as exc:
        print("WRITE SMOKE TEST", file=sys.stderr)
        print(f"- verification: FAIL", file=sys.stderr)
        print(f"- error: {exc}", file=sys.stderr)
        return 2
    except Exception as exc:
        print("WRITE SMOKE TEST", file=sys.stderr)
        print(f"- verification: FAIL", file=sys.stderr)
        print(f"- error: unexpected failure: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
