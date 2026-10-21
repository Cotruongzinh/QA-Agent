"""Read-only structural inspection of the project's QA Excel template.

This script intentionally has no save or output-workbook operation. It prints a
dry-run report to stdout and closes the workbook without modifying it.
"""

from __future__ import annotations

import argparse
import re
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable, Sequence

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet


DEFAULT_TEMPLATE = Path(r"C:\QA-Agent\templates\QA_Report_Template.xlsx")

TESTCASE_HEADERS = (
    "TC_ID",
    "Module",
    "Test Title",
    "Description/Test Summary",
    "Test Steps",
    "Test Data",
    "Expected Result",
    "Actual Result",
    "Log/Screenshot/Video",
    "Result",
    "Phản hồi",
    "FEEDBACK",
)

BUG_HEADERS = (
    "Bug ID",
    "Module",
    "Description",
    "Steps",
    "Expected Result",
    "Actual Result",
    "Severity",
    "Priority",
    "Status",
    "Log/Screenshot/Video",
)

SUMMARY_ANCHORS = {
    "summary",
    "qa summary",
    "test summary",
    "test result summary",
    "tổng quan",
    "tổng hợp",
}

KPI_TERMS = (
    "total",
    "testcase",
    "test case",
    "executed",
    "pass",
    "fail",
    "blocked",
    "not run",
    "rate",
    "bug",
    "defect",
    "tổng",
    "đạt",
    "không đạt",
)


def normalized(value: object) -> str:
    if value is None:
        return ""
    return " ".join(str(value).strip().casefold().split())


def display_value(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def iter_populated_cells(ws: Worksheet) -> Iterable[object]:
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None:
                yield cell


def row_values(ws: Worksheet, row_number: int) -> list[object]:
    return [ws.cell(row=row_number, column=column).value for column in range(1, ws.max_column + 1)]


def header_rows(ws: Worksheet, anchor: str) -> list[int]:
    target = normalized(anchor)
    matches: list[int] = []
    for row_number in range(1, ws.max_row + 1):
        if any(normalized(value) == target for value in row_values(ws, row_number)):
            matches.append(row_number)
    return matches


def header_score(ws: Worksheet, row_number: int, expected_headers: Sequence[str]) -> int:
    actual = {normalized(value) for value in row_values(ws, row_number) if value is not None}
    expected = {normalized(value) for value in expected_headers}
    return len(actual & expected)


@dataclass(frozen=True)
class DataBlock:
    worksheet: Worksheet
    header_row: int
    anchor_column: int
    column_map: dict[str, str]
    data_start_row: int
    data_end_row: int | None
    record_rows: tuple[int, ...]


def make_column_map(ws: Worksheet, header_row: int) -> dict[str, str]:
    mapping: dict[str, str] = {}
    seen: dict[str, int] = {}
    for column in range(1, ws.max_column + 1):
        label = display_value(ws.cell(header_row, column).value)
        if not label:
            continue
        seen[label] = seen.get(label, 0) + 1
        key = label if seen[label] == 1 else f"{label} (#{seen[label]})"
        mapping[key] = get_column_letter(column)
    return mapping


def detect_data_block(
    workbook: object,
    anchor: str,
    expected_headers: Sequence[str],
    title_hint: str,
) -> DataBlock | None:
    candidates: list[tuple[int, int, int, int, Worksheet, int]] = []
    for sheet_index, ws in enumerate(workbook.worksheets):
        occurrences = header_rows(ws, anchor)
        if not occurrences:
            continue
        current_header = occurrences[-1]
        score = header_score(ws, current_header, expected_headers)
        title_match = int(normalized(title_hint) in normalized(ws.title).replace(" ", ""))
        candidates.append((score, title_match, current_header, sheet_index, ws, current_header))

    if not candidates:
        return None

    _, _, _, _, ws, current_header = max(candidates, key=lambda item: item[:4])
    anchor_column = next(
        column
        for column in range(1, ws.max_column + 1)
        if normalized(ws.cell(current_header, column).value) == normalized(anchor)
    )
    record_rows = tuple(
        row
        for row in range(current_header + 1, ws.max_row + 1)
        if display_value(ws.cell(row, anchor_column).value)
        and normalized(ws.cell(row, anchor_column).value) != normalized(anchor)
    )
    return DataBlock(
        worksheet=ws,
        header_row=current_header,
        anchor_column=anchor_column,
        column_map=make_column_map(ws, current_header),
        data_start_row=current_header + 1,
        data_end_row=max(record_rows) if record_rows else None,
        record_rows=record_rows,
    )


def find_column(block: DataBlock, names: Sequence[str]) -> int | None:
    wanted = {normalized(name) for name in names}
    for column in range(1, block.worksheet.max_column + 1):
        if normalized(block.worksheet.cell(block.header_row, column).value) in wanted:
            return column
    return None


def count_results(block: DataBlock) -> tuple[int | None, int | None]:
    result_column = find_column(block, ("Result", "Status", "Kết quả"))
    if result_column is None:
        return None, None
    values = [normalized(block.worksheet.cell(row, result_column).value) for row in block.record_rows]
    pass_count = sum(value in {"pass", "passed", "đạt"} for value in values)
    fail_count = sum(value in {"fail", "failed", "không đạt"} for value in values)
    return pass_count, fail_count


def is_formula(value: object) -> bool:
    return isinstance(value, str) and value.startswith("=")


def summary_signal_rows(ws: Worksheet) -> list[int]:
    signals: list[int] = []
    for row_number in range(1, ws.max_row + 1):
        values = [normalized(value) for value in row_values(ws, row_number) if value is not None]
        anchor_found = any(value in SUMMARY_ANCHORS for value in values)
        kpi_hits = sum(any(term in value for term in KPI_TERMS) for value in values)
        if anchor_found or kpi_hits >= 1:
            signals.append(row_number)
    return signals


def cluster_rows(rows: Sequence[int], max_gap: int = 6) -> list[list[int]]:
    clusters: list[list[int]] = []
    for row in sorted(set(rows)):
        if not clusters or row - clusters[-1][-1] > max_gap:
            clusters.append([row])
        else:
            clusters[-1].append(row)
    return clusters


@dataclass(frozen=True)
class SummaryBlock:
    worksheet: Worksheet
    start_row: int
    end_row: int
    min_column: int
    max_column: int
    formulas: tuple[tuple[str, str, str], ...]


def detect_summary_block(workbook: object) -> SummaryBlock | None:
    candidates: list[tuple[int, int, int, int, Worksheet, list[list[int]]]] = []
    for sheet_index, ws in enumerate(workbook.worksheets):
        signals = summary_signal_rows(ws)
        clusters = cluster_rows(signals)
        if not clusters:
            continue
        title_bonus = 10 if "summary" in normalized(ws.title) else 0
        anchor_count = sum(
            normalized(cell.value) in SUMMARY_ANCHORS for cell in iter_populated_cells(ws)
        )
        latest_start = clusters[-1][0]
        candidates.append(
            (title_bonus + anchor_count, len(clusters[-1]), latest_start, sheet_index, ws, clusters)
        )

    if not candidates:
        return None

    _, _, _, _, ws, clusters = max(candidates, key=lambda item: item[:4])
    latest_cluster = clusters[-1]
    start_row = latest_cluster[0]

    populated_rows = [
        row
        for row in range(start_row, ws.max_row + 1)
        if any(ws.cell(row, column).value is not None for column in range(1, ws.max_column + 1))
    ]
    end_row = max(populated_rows) if populated_rows else start_row

    populated_columns = [
        column
        for row in range(start_row, end_row + 1)
        for column in range(1, ws.max_column + 1)
        if ws.cell(row, column).value is not None
    ]
    min_column = min(populated_columns) if populated_columns else 1
    max_column = max(populated_columns) if populated_columns else 1

    formulas: list[tuple[str, str, str]] = []
    fallback_formulas: list[tuple[str, str, str]] = []
    for row in range(start_row, end_row + 1):
        row_labels = [
            display_value(ws.cell(row, column).value)
            for column in range(min_column, max_column + 1)
            if ws.cell(row, column).value is not None
            and not is_formula(ws.cell(row, column).value)
        ]
        label_text = " | ".join(row_labels)
        label_norm = normalized(label_text)
        for column in range(min_column, max_column + 1):
            cell = ws.cell(row, column)
            if not is_formula(cell.value):
                continue
            item = (cell.coordinate, label_text or "(no label on row)", str(cell.value))
            fallback_formulas.append(item)
            if any(term in label_norm for term in KPI_TERMS):
                formulas.append(item)

    return SummaryBlock(
        worksheet=ws,
        start_row=start_row,
        end_row=end_row,
        min_column=min_column,
        max_column=max_column,
        formulas=tuple(formulas or fallback_formulas),
    )


def formula_range_references(formula: str, sheet_name: str) -> list[tuple[int, int, str]]:
    escaped_name = re.escape(sheet_name.replace("'", "''"))
    sheet_token = rf"(?:'{escaped_name}'|{re.escape(sheet_name)})!"
    pattern = re.compile(
        sheet_token
        + r"\$?[A-Z]{1,3}\$?(\d+)\s*:\s*\$?[A-Z]{1,3}\$?(\d+)",
        flags=re.IGNORECASE,
    )
    references: list[tuple[int, int, str]] = []
    for match in pattern.finditer(formula):
        start, end = sorted((int(match.group(1)), int(match.group(2))))
        references.append((start, end, match.group(0)))
    return references


def suspicious_references(
    summary: SummaryBlock | None,
    data_blocks: Sequence[tuple[str, DataBlock | None]],
) -> list[str]:
    if summary is None:
        return []
    findings: list[str] = []
    for coordinate, _, formula in summary.formulas:
        for block_name, block in data_blocks:
            if block is None or block.data_end_row is None:
                continue
            expected_start = block.data_start_row
            expected_end = block.data_end_row
            for actual_start, actual_end, reference in formula_range_references(
                formula, block.worksheet.title
            ):
                reasons: list[str] = []
                if actual_end < expected_start or actual_start > expected_end:
                    reasons.append("range is outside the current block")
                else:
                    if actual_start < expected_start:
                        reasons.append("range starts in an older block")
                    elif actual_start > expected_start:
                        reasons.append("range skips the start of the current block")
                    if actual_end < expected_end:
                        reasons.append("range ends before current data ends")
                if reasons:
                    findings.append(
                        f"{summary.worksheet.title}!{coordinate}: {reference} ({block_name}: "
                        + "; ".join(reasons)
                        + ")"
                    )
    return findings


def conditional_formatting_details(ws: Worksheet) -> tuple[int, int, list[str]]:
    ranges: list[str] = []
    rule_count = 0
    for item in ws.conditional_formatting:
        ranges.append(str(getattr(item, "sqref", item)))
        try:
            rule_count += len(ws.conditional_formatting[item])
        except (KeyError, TypeError):
            pass
    return len(ranges), rule_count, ranges


def image_anchor(image: object) -> str:
    anchor = getattr(image, "anchor", None)
    marker = getattr(anchor, "_from", None)
    if marker is None:
        return "unknown"
    return f"{get_column_letter(marker.col + 1)}{marker.row + 1}"


def image_details(ws: Worksheet) -> list[str]:
    details: list[str] = []
    for index, image in enumerate(getattr(ws, "_images", ()), start=1):
        image_format = getattr(image, "format", None) or "unknown format"
        width = getattr(image, "width", "?")
        height = getattr(image, "height", "?")
        details.append(
            f"image {index}: anchor={image_anchor(image)}, format={image_format}, size={width}x{height}"
        )
    return details


def formula_count(ws: Worksheet) -> int:
    return sum(is_formula(cell.value) for cell in iter_populated_cells(ws))


def styled_cell_count(ws: Worksheet) -> int:
    return sum(bool(cell.has_style) for row in ws.iter_rows() for cell in row)


def format_rows(rows: Sequence[int]) -> str:
    return ", ".join(str(row) for row in rows) if rows else "none"


def print_column_map(mapping: dict[str, str]) -> None:
    if not mapping:
        print("  column mapping: not found")
        return
    print("  column mapping:")
    for name, letter in mapping.items():
        print(f"    {name} -> {letter}")


def print_data_block(title: str, block: DataBlock | None, include_results: bool = False) -> None:
    print(f"\n{title}")
    if block is None:
        print("  current block: not found")
        return
    print(f"  sheet: {block.worksheet.title}")
    print(f"  current header row: {block.header_row}")
    print_column_map(block.column_map)
    print(f"  current data start row: {block.data_start_row}")
    print(f"  current data end row: {block.data_end_row if block.data_end_row else 'none'}")
    noun = "testcases" if title == "TESTCASE" else "bugs"
    print(f"  current {noun}: {len(block.record_rows)}")
    if include_results:
        pass_count, fail_count = count_results(block)
        print(f"  Pass: {pass_count if pass_count is not None else 'not readable'}")
        print(f"  Fail: {fail_count if fail_count is not None else 'not readable'}")


def print_summary(
    summary: SummaryBlock | None,
    testcase: DataBlock | None,
    bug: DataBlock | None,
) -> None:
    print("\nSUMMARY")
    if summary is None:
        print("  current block: not found")
        return
    location = (
        f"{summary.worksheet.title}!"
        f"{get_column_letter(summary.min_column)}{summary.start_row}:"
        f"{get_column_letter(summary.max_column)}{summary.end_row}"
    )
    print(f"  current block: {location}")
    print("  KPI formulas:")
    if summary.formulas:
        for coordinate, label, formula in summary.formulas:
            print(f"    {coordinate} [{label}]: {formula}")
    else:
        print("    none found")

    findings = suspicious_references(
        summary, (("TESTCASE", testcase), ("NOTEBUG", bug))
    )
    print("  possible wrong-block references:")
    if findings:
        for finding in findings:
            print(f"    WARNING: {finding}")
    else:
        print("    none detected by row-range heuristic")


def print_structure(workbook: object) -> None:
    print("\nSTRUCTURE")
    total_images = 0
    for ws in workbook.worksheets:
        merges = [str(cell_range) for cell_range in ws.merged_cells.ranges]
        images = image_details(ws)
        total_images += len(images)
        cf_range_count, cf_rule_count, cf_ranges = conditional_formatting_details(ws)
        hidden_rows = sorted(index for index, dim in ws.row_dimensions.items() if dim.hidden)
        hidden_columns = sorted(key for key, dim in ws.column_dimensions.items() if dim.hidden)
        print(f"  [{ws.title}]")
        print(f"    merged cells ({len(merges)}): {', '.join(merges) if merges else 'none'}")
        print(f"    formula count: {formula_count(ws)}")
        print(f"    styled cell count: {styled_cell_count(ws)}")
        print(f"    image count: {len(images)}")
        for detail in images:
            print(f"      {detail}")
        print(
            f"    conditional formatting: {cf_range_count} range(s), "
            f"{cf_rule_count} rule(s)"
        )
        if cf_ranges:
            print(f"      ranges: {', '.join(cf_ranges)}")
        print(f"    freeze panes: {ws.freeze_panes or 'none'}")
        print(f"    hidden rows: {format_rows(hidden_rows)}")
        print(f"    hidden columns: {', '.join(hidden_columns) if hidden_columns else 'none'}")

    print("\nEMBEDDED IMAGE CHECK")
    print(f"  embedded screenshot/image detected: {'YES' if total_images else 'NO'}")
    print(f"  total embedded images: {total_images}")


def inspect(template_path: Path) -> None:
    if not template_path.is_file():
        raise FileNotFoundError(f"Template not found: {template_path}")
    if template_path.suffix.casefold() not in {".xlsx", ".xlsm"}:
        raise ValueError("Expected an .xlsx or .xlsm workbook")

    print("DRY RUN ONLY - workbook will not be saved")
    workbook = load_workbook(
        filename=template_path,
        read_only=False,
        data_only=False,
        keep_vba=template_path.suffix.casefold() == ".xlsm",
    )
    try:
        print("\nWORKBOOK")
        print(f"  path: {template_path.resolve()}")
        print(f"  sheet names: {', '.join(workbook.sheetnames)}")
        print("  dimensions:")
        for ws in workbook.worksheets:
            print(f"    {ws.title}: {ws.max_row} row(s) x {ws.max_column} column(s) [{ws.dimensions}]")

        testcase = detect_data_block(workbook, "TC_ID", TESTCASE_HEADERS, "testcase")
        bug = detect_data_block(workbook, "Bug ID", BUG_HEADERS, "notebug")
        summary = detect_summary_block(workbook)

        print_data_block("TESTCASE", testcase, include_results=True)
        print_data_block("NOTEBUG", bug)
        print_summary(summary, testcase, bug)
        print_structure(workbook)
    finally:
        workbook.close()


def parse_args(argv: Sequence[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Inspect the QA Excel template without saving or modifying it."
    )
    parser.add_argument(
        "template",
        nargs="?",
        type=Path,
        default=DEFAULT_TEMPLATE,
        help=f"Workbook to inspect (default: {DEFAULT_TEMPLATE})",
    )
    return parser.parse_args(argv)


def main(argv: Sequence[str] | None = None) -> int:
    args = parse_args(argv)
    try:
        inspect(args.template)
    except (FileNotFoundError, ValueError) as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 2
    except Exception as exc:
        print(f"ERROR: inspection failed: {exc}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
